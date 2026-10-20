#!/usr/bin/env python3
"""Create a deterministic, read-only profile for CSV, TSV, JSON, or XLSX data."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

MISSING = {"", "na", "n/a", "null", "none", "nan", "-"}
SENSITIVE = re.compile(r"(?:^id$|name|姓名|student.?id|学号|patient.?id|subject.?id|participant.?id|record.?id|患者|id.?card|身份证|phone|mobile|电话|email|邮箱|address|住址|password|token|secret|credential)", re.I)
SPATIAL = re.compile(r"(?:latitude|longitude|\blat\b|\blon\b|\blng\b|province|city|county|district|经度|纬度|省|市|县|区)", re.I)
UNIT = re.compile(r"(?:\(([^()]+)\)|（([^（）]+)）|\[([^\[\]]+)\])\s*$")
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m", "%Y/%m")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return f"sha256:{digest.hexdigest()}"


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def missing(value: Any) -> bool:
    return clean(value).lower() in MISSING


def parse_number(value: str) -> float | None:
    try:
        number = float(value.replace(",", ""))
        return number if math.isfinite(number) else None
    except ValueError:
        return None


def parse_date(value: str) -> datetime | None:
    candidate = value.strip()
    if not candidate:
        return None
    try:
        return datetime.fromisoformat(candidate.replace("Z", "+00:00"))
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(candidate, fmt)
        except ValueError:
            continue
    return None


def normalize_rows(rows: Iterable[dict[str, Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    records = list(rows)
    fields: list[str] = []
    seen: set[str] = set()
    for row in records:
        for key in row:
            name = clean(key)
            if name and name not in seen:
                seen.add(name)
                fields.append(name)
    return fields, [{field: row.get(field) for field in fields} for row in records]


def read_delimited(path: Path, delimiter: str) -> dict[str, list[dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {path.stem: list(csv.DictReader(handle, delimiter=delimiter))}


def read_json(path: Path) -> dict[str, list[dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)
    if isinstance(payload, dict):
        if isinstance(payload.get("records"), list):
            payload = payload["records"]
        elif all(isinstance(value, list) for value in payload.values()):
            tables = {}
            for key, value in payload.items():
                if not all(isinstance(row, dict) for row in value):
                    raise ValueError(f"JSON table '{key}' must be a list of objects")
                tables[str(key)] = value
            return tables
        else:
            payload = [payload]
    if not isinstance(payload, list) or not all(isinstance(row, dict) for row in payload):
        raise ValueError("JSON must be an object, a list of objects, or an object containing record lists")
    return {path.stem: payload}


def read_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as error:
        raise RuntimeError("XLSX profiling requires openpyxl in the active Python runtime") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: dict[str, list[dict[str, Any]]] = {}
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = [clean(value) or f"column_{index + 1}" for index, value in enumerate(next(iterator))]
        except StopIteration:
            tables[sheet.title] = []
            continue
        tables[sheet.title] = [dict(zip(header, row)) for row in iterator]
    return tables


def load_tables(path: Path) -> dict[str, list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_delimited(path, ",")
    if suffix == ".tsv":
        return read_delimited(path, "\t")
    if suffix == ".json":
        return read_json(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    raise ValueError("supported dataset formats are CSV, TSV, JSON, and XLSX")


def profile_column(name: str, values: list[Any], row_count: int) -> dict[str, Any]:
    present = [clean(value) for value in values if not missing(value)]
    missing_count = row_count - len(present)
    distinct = sorted(set(present))
    result: dict[str, Any] = {
        "name": name,
        "missingCount": missing_count,
        "missingRate": round(missing_count / row_count, 6) if row_count else 0,
        "distinctCount": len(distinct),
        "sensitiveCandidate": bool(SENSITIVE.search(name)),
        "spatialCandidate": bool(SPATIAL.search(name)),
    }
    unit_match = UNIT.search(name)
    if unit_match:
        result["unitCandidate"] = next(group for group in unit_match.groups() if group)
    if not present:
        result["inferredType"] = "empty"
        return result
    lowered = {value.lower() for value in present}
    if lowered <= {"true", "false", "yes", "no", "0", "1"}:
        result["inferredType"] = "boolean"
        return result
    numbers = [parse_number(value) for value in present]
    if all(value is not None for value in numbers):
        numeric = [value for value in numbers if value is not None]
        result.update({"inferredType": "number", "minimum": min(numeric), "maximum": max(numeric)})
        return result
    dates = [parse_date(value) for value in present]
    if all(value is not None for value in dates):
        parsed = [value for value in dates if value is not None]
        result.update({"inferredType": "datetime", "minimum": min(parsed).isoformat(), "maximum": max(parsed).isoformat()})
        return result
    threshold = min(50, max(20, int(math.sqrt(max(row_count, 1)) * 2)))
    result["inferredType"] = "categorical" if len(distinct) <= threshold else "text"
    if result["inferredType"] == "categorical" and not result["sensitiveCandidate"]:
        result["sampleValues"] = distinct[:10]
    return result


def profile_table(name: str, raw_rows: list[dict[str, Any]]) -> dict[str, Any]:
    fields, rows = normalize_rows(raw_rows)
    row_count = len(rows)
    signatures = [tuple(clean(row.get(field)) for field in fields) for row in rows]
    duplicate_count = row_count - len(set(signatures))
    columns = [profile_column(field, [row.get(field) for row in rows], row_count) for field in fields]
    missing_cells = sum(column["missingCount"] for column in columns)
    total_cells = row_count * len(fields)
    primary_key_candidates = [
        column["name"] for column in columns
        if column["missingCount"] == 0 and column["distinctCount"] == row_count and row_count > 0
    ]
    time_columns = [column for column in columns if column["inferredType"] == "datetime"]
    return {
        "name": name,
        "rowCount": row_count,
        "columnCount": len(fields),
        "duplicateRows": duplicate_count,
        "missingCells": missing_cells,
        "missingRate": round(missing_cells / total_cells, 6) if total_cells else 0,
        "observationGrain": "requires interpretation from source context",
        "primaryKeyCandidates": primary_key_candidates,
        "timeCoverage": [{"field": column["name"], "minimum": column["minimum"], "maximum": column["maximum"]} for column in time_columns],
        "spatialFieldCandidates": [column["name"] for column in columns if column["spatialCandidate"]],
        "sensitiveFieldCandidates": [column["name"] for column in columns if column["sensitiveCandidate"]],
        "columns": columns,
    }


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: profile_dataset.py <dataset.csv|tsv|json|xlsx> [output.json]", file=sys.stderr)
        return 2
    source = Path(sys.argv[1]).expanduser().resolve()
    if not source.is_file():
        print(f"[dataset-profiler] source is not a file: {source}", file=sys.stderr)
        return 2
    try:
        tables = load_tables(source)
        report = {
            "source": str(source),
            "sourceHash": sha256(source),
            "format": source.suffix.lower().lstrip("."),
            "tables": [profile_table(name, rows) for name, rows in tables.items()],
        }
    except Exception as error:  # Present a deterministic, readable dependency/format error.
        print(f"[dataset-profiler] {error}", file=sys.stderr)
        return 1
    encoded = json.dumps(report, ensure_ascii=False, indent=2)
    if len(sys.argv) > 2:
        output = Path(sys.argv[2]).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(encoded + "\n", encoding="utf-8")
    else:
        print(encoded)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
